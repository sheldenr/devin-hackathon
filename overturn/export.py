"""Export appeal letter to .docx (fully local)."""
from io import BytesIO

from docx import Document


def letter_to_docx(letter_text: str, title: str = "Appeal Letter") -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for para in letter_text.split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
